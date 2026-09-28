"""Regenerate template/contacts_template.xlsx and the fictional files in examples/.

All sample people and organizations are made up.
Run: python scripts/make_template.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
COLUMNS = ["Name", "Company", "School", "Role", "Connected Through",
           "LinkedIn URL", "Status", "Tags", "Notes"]
WIDTHS = [22, 24, 26, 26, 20, 36, 14, 22, 40]
STATUSES = ["Met", "To Reach Out", "Contacted", "Follow Up", "Referral"]

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def styled_sheet(wb: Workbook, rows: list[list[str]], targets: list[list[str]] = ()):
    ws = wb.active
    ws.title = "Contacts"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22
    dv = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("G2:G2000")
    for r in rows:
        ws.append(r)
    ws.auto_filter.ref = f"A1:I{max(2, ws.max_row)}"

    help_ws = wb.create_sheet("How to use")
    help_ws.column_dimensions["A"].width = 22
    help_ws.column_dimensions["B"].width = 90
    notes = [
        ("Name", "Required. Full name, spelled the same way everywhere."),
        ("Company", "Current company. 3+ direct connections at one company become a group."),
        ("School", "School they attended. 3+ at one school become a group. 'BYU' and "
                   "'Brigham Young University' count as the same."),
        ("Role", "Job title (optional)."),
        ("Connected Through", "Leave BLANK if you know them directly. If someone introduced you "
                              "(or you know OF them through someone), put that person's name here "
                              "and they appear as a 2nd-degree connection."),
        ("LinkedIn URL", "Paste their profile link. Double-click them on the map to open it."),
        ("Status", "Met / To Reach Out / Contacted / Follow Up / Referral."),
        ("Tags", "Optional, comma-separated (e.g. 'MCA, Delta recruiting'). "
                 "Group by tags with: --group-by company,school,tags"),
        ("Notes", "Anything you want to remember."),
        ("", ""),
        ("Targets tab", "List companies you want to work at (one per row). Targets are outlined in "
                        "red on the map; ones with no connections show as hollow red bubbles, and "
                        "build_map.py writes a 'who can intro me?' report for each."),
        ("LinkedIn export", "LinkedIn > Settings > Data privacy > Get a copy of your data > "
                            "Connections. Pass the Connections.csv with --linkedin. Rows here "
                            "with the same name override the export (so you can add schools)."),
    ]
    for k, v in notes:
        help_ws.append([k, v])
        help_ws.cell(help_ws.max_row, 1).font = Font(bold=True)
        help_ws.cell(help_ws.max_row, 2).alignment = Alignment(wrap_text=True)

    target_ws = wb.create_sheet("Targets", 1)
    target_ws.append(["Company", "Notes"])
    for cell in target_ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    target_ws.column_dimensions["A"].width = 30
    target_ws.column_dimensions["B"].width = 50
    target_ws.freeze_panes = "A2"
    for r in targets:
        target_ws.append(r)
    return wb


SAMPLE = [
    # Name, Company, School, Role, Connected Through, URL, Status, Tags, Notes
    ["Jordan Lee", "Northwind Consulting", "Riverbend University", "Senior Consultant", "", "", "Met", "Consulting club", "Case interview tips"],
    ["Priya Shah", "Northwind Consulting", "Lakeview State University", "Manager", "", "", "Contacted", "", ""],
    ["Marcus Bell", "Northwind Consulting", "Riverbend University", "Analyst", "", "", "Met", "Consulting club", ""],
    ["Hana Kim", "Northwind Consulting, Inc.", "", "Partner", "", "", "To Reach Out", "", "Leads the strategy practice"],
    ["Sofia Alvarez", "Contoso Games", "Riverbend University", "Product Manager", "", "", "Met", "Gaming", "Owns the marketplace roadmap"],
    ["Ethan Brooks", "Contoso Games", "", "Data Analyst", "", "", "Follow Up", "Gaming", ""],
    ["Lena Novak", "Contoso Games", "Harbor Tech Institute", "Producer", "", "", "To Reach Out", "Gaming", ""],
    ["Noah Carter", "Summit Airlines", "Riverbend University", "Commercial Strategy Analyst", "", "", "Met", "", "Info session speaker"],
    ["Grace Owens", "Summit Airlines", "", "Recruiter", "", "", "Contacted", "", ""],
    ["Daniel Ortiz", "Brightline Bank", "Riverbend University", "Strategy Associate", "", "", "Met", "", ""],
    ["Mia Chen", "Fieldstone Capital", "Lakeview State University", "Partner", "", "", "Met", "", ""],
    ["Owen Price", "Fieldstone Capital", "Lakeview State University", "Associate", "", "", "To Reach Out", "", ""],
    ["Ava Thompson", "Keystone Health", "Lakeview State University", "Operations Lead", "", "", "Met", "", ""],
    ["Liam Walsh", "Independent", "", "Founder", "", "", "Met", "", "Old roommate"],
    # 2nd-degree
    ["Rachel Green", "Northwind Consulting", "", "Principal", "Jordan Lee", "", "To Reach Out", "", "Jordan offered an intro"],
    ["Tom Nguyen", "Contoso Games", "", "Director of Product", "Sofia Alvarez", "", "To Reach Out", "Gaming", ""],
    ["Isla Moore", "Contoso Games", "", "UX Researcher", "Sofia Alvarez", "", "", "Gaming", ""],
    ["Victor Hale", "Summit Airlines", "", "VP Network Planning", "Noah Carter", "", "", "", ""],
    ["Chloe Park", "Harbor Ventures", "", "Investor", "Mia Chen", "", "", "", ""],
    ["Ben Foster", "Riverbend University", "", "Career Coach", "Daniel Ortiz", "", "", "", ""],
    ["Zoe Adams", "Pinecrest Labs", "", "CEO", "Liam Walsh", "", "", "", ""],
    ["Sam Rivera", "Pinecrest Labs", "", "Head of Growth", "Zoe Adams", "", "", "", "3rd-degree example"],
]

TARGETS = [
    ["Northwind Consulting", "Strategy practice"],
    ["Summit Airlines", ""],
    ["Pinecrest Labs", "Reachable only through Liam"],
    ["Granite Peak Partners", "No one yet"],
    ["Harborview Media", "No one yet"],
]

LINKEDIN = """Notes:
"When exporting your connection data, you may notice that some of the email addresses are missing."

First Name,Last Name,URL,Email Address,Company,Position,Connected On
Jordan,Lee,https://example.com/in/jordan-lee,,Northwind Consulting,Senior Consultant,12 Mar 2026
Priya,Shah,https://example.com/in/priya-shah,,Northwind Consulting,Manager,02 Feb 2026
Sofia,Alvarez,https://example.com/in/sofia-alvarez,,Contoso Games,Product Manager,21 Aug 2026
Kai,Morgan,https://example.com/in/kai-morgan,,Brightline Bank,Analyst,15 Jan 2026
Elena,Rossi,https://example.com/in/elena-rossi,,Brightline Bank,Risk Associate,03 May 2026
Theo,Grant,https://example.com/in/theo-grant,,Brightline Bank,Product Owner,09 Jun 2026
Nora,Singh,https://example.com/in/nora-singh,,Keystone Health,Strategy Manager,30 Jul 2026
"""


def main():
    (ROOT / "template").mkdir(exist_ok=True)
    (ROOT / "examples").mkdir(exist_ok=True)
    styled_sheet(Workbook(), []).save(ROOT / "template" / "contacts_template.xlsx")
    styled_sheet(Workbook(), SAMPLE, TARGETS).save(ROOT / "examples" / "sample_contacts.xlsx")
    (ROOT / "examples" / "sample_linkedin_connections.csv").write_text(LINKEDIN, encoding="utf-8")
    print("Wrote template/contacts_template.xlsx and examples/")


if __name__ == "__main__":
    main()
